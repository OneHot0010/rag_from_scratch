"""Excel 加载器(.xlsx/.xlsm,基于 openpyxl)。

阶段2.1 增强(P0 正确性):
  - 合并单元格填充:openpyxl 只在左上角返回值,这里把合并区域回填,
    避免分类/表头列大面积丢值;
  - 可配置表头:header_row 指定表头起始行、header_depth 支持多级表头合并,
    应对"标题行/空行/多行表头"等真实报表结构;
  - 类型规范化:日期→ISO、数字去 .0、bool 显式化(见 _tabular.normalize_value);
  - 跳过空行;每行注入 sheet 名作为语义锚点(缓解多 sheet 同名列歧义);
  - 逐 sheet 容错:单个 sheet 解析失败不阻断其余 sheet。

需要合并单元格信息时 openpyxl 的 read_only 模式不暴露 merged_cells,
故按需切换:fill_merged=True 用普通模式,否则用 read_only 省内存。
"""
from __future__ import annotations

import logging
from pathlib import Path
from typing import List, Optional

from ..models import Document
from .base import DataLoader
from ._tabular import (
    clean_headers,
    merge_header_rows,
    row_to_text,
    is_blank_row,
)

logger = logging.getLogger(__name__)

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None


class ExcelLoader(DataLoader):
    extensions = (".xlsx", ".xlsm")

    def __init__(self, header_row: int = 1, header_depth: int = 1,
                 fill_merged: bool = True, inject_sheet_name: bool = True) -> None:
        """
        Args:
            header_row:   表头起始行(1-based),之前的行(标题/说明)会被跳过。
            header_depth: 表头占用的行数,>1 时按多级表头合并列名。
            fill_merged:  是否回填合并单元格的值。
            inject_sheet_name: 是否把 sheet 名作为语义锚点前置到每行文本。
        """
        self.header_row = max(1, header_row)
        self.header_depth = max(1, header_depth)
        self.fill_merged = fill_merged
        self.inject_sheet_name = inject_sheet_name

    def load(self, path: str | Path) -> List[Document]:
        if load_workbook is None:
            raise ImportError("需要 openpyxl,请先 `pip install openpyxl`。")
        p = self._ensure_exists(path)
        if p.suffix.lower() == ".xls":  # 双保险:老格式 openpyxl 不支持
            raise ValueError("openpyxl 不支持旧版 .xls,请转存为 .xlsx 或接入 xlrd。")

        read_only = not self.fill_merged
        wb = load_workbook(filename=str(p), read_only=read_only, data_only=True)
        docs: List[Document] = []
        try:
            for ws in wb.worksheets:
                try:
                    docs.extend(self._load_sheet(ws, source=str(p)))
                except Exception as e:  # 逐 sheet 隔离容错
                    logger.warning("解析 sheet 失败 [%s!%s]: %s", p.name, ws.title, e)
        finally:
            wb.close()
        return docs

    def _read_matrix(self, ws) -> List[list]:
        """读出二维值矩阵,并按需回填合并单元格。"""
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if self.fill_merged and getattr(ws, "merged_cells", None):
            for rng in list(ws.merged_cells.ranges):
                top = ws.cell(row=rng.min_row, column=rng.min_col).value
                for r in range(rng.min_row, rng.max_row + 1):
                    for c in range(rng.min_col, rng.max_col + 1):
                        ri, ci = r - 1, c - 1
                        if 0 <= ri < len(rows) and 0 <= ci < len(rows[ri]):
                            if rows[ri][ci] is None:
                                rows[ri][ci] = top
        return rows

    def _load_sheet(self, ws, source: str) -> List[Document]:
        matrix = self._read_matrix(ws)
        if not matrix:
            return []

        h_start = self.header_row - 1
        h_end = h_start + self.header_depth
        if h_start >= len(matrix):
            return []

        if self.header_depth > 1:
            headers = merge_header_rows(matrix[h_start:h_end])
        else:
            headers = clean_headers(matrix[h_start])

        sheet_name: Optional[str] = ws.title if self.inject_sheet_name else None
        docs: List[Document] = []
        for offset, values in enumerate(matrix[h_end:]):
            row_idx = h_end + offset + 1  # 1-based 真实行号
            if is_blank_row(values):
                continue
            content = row_to_text(headers, values, sheet=sheet_name)
            doc = Document(
                content=content,
                metadata={"source": source, "format": "excel",
                          "sheet": ws.title, "row": row_idx},
            )
            if doc:
                docs.append(doc)
        return docs
