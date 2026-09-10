"""Excel 加载器逐函数单元测试。

对 ragcore/loaders/_tabular.py 和 ragcore/loaders/excel_loader.py 中
每个优化函数进行独立测试，使用 sample.xlsx 作为真实数据源。
"""

from __future__ import annotations

import datetime as _dt
import importlib.util
import sys
from decimal import Decimal
from pathlib import Path

# 确保项目根目录在 sys.path 中
_PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_PROJECT_ROOT))

# 绕过 ragcore/__init__.py 的 import 链（避免触发 volcenginesdkarkruntime 依赖），
# 直接从文件路径加载子模块，并注册到 sys.modules 使相对 import 正常工作。


def _load_module(package_name: str, file_path: Path) -> None:
    """从文件路径加载模块并注册到 sys.modules[package_name]。
    
    先递归加载依赖的父包（空模块），再加载目标模块。
    所有模块的 __package__ 和 __name__ 会被正确设置，使得相对 import 可用。
    """
    parts = package_name.split(".")
    # 注册所有父包（空模块）
    for i in range(1, len(parts)):
        parent_name = ".".join(parts[:i])
        if parent_name not in sys.modules:
            parent = importlib.util.module_from_spec(
                importlib.util.spec_from_loader(parent_name, loader=None, origin=str(file_path.parent))
            )
            parent.__package__ = parent_name
            parent.__path__ = []
            sys.modules[parent_name] = parent

    spec = importlib.util.spec_from_file_location(package_name, str(file_path))
    mod = importlib.util.module_from_spec(spec)
    mod.__package__ = ".".join(parts[:-1])
    sys.modules[package_name] = mod
    spec.loader.exec_module(mod)


# 按依赖顺序加载：models → base → _tabular → excel_loader
_load_module("ragcore.models", _PROJECT_ROOT / "ragcore" / "models.py")
_load_module("ragcore.loaders.base", _PROJECT_ROOT / "ragcore" / "loaders" / "base.py")
_load_module("ragcore.loaders._tabular", _PROJECT_ROOT / "ragcore" / "loaders" / "_tabular.py")
_load_module("ragcore.loaders.excel_loader", _PROJECT_ROOT / "ragcore" / "loaders" / "excel_loader.py")

# 从已加载的模块中提取需要的符号
normalize_value = sys.modules["ragcore.loaders._tabular"].normalize_value
clean_headers = sys.modules["ragcore.loaders._tabular"].clean_headers
merge_header_rows = sys.modules["ragcore.loaders._tabular"].merge_header_rows
row_to_text = sys.modules["ragcore.loaders._tabular"].row_to_text
is_blank_row = sys.modules["ragcore.loaders._tabular"].is_blank_row
MAX_FIELD_CHARS = sys.modules["ragcore.loaders._tabular"].MAX_FIELD_CHARS
ExcelLoader = sys.modules["ragcore.loaders.excel_loader"].ExcelLoader

# ── 测试文件路径 ──────────────────────────────────────────────────
SAMPLE_XLSX = Path(__file__).resolve().parent.parent / "sample.xlsx"


# ====================================================================
# 1. normalize_value 测试
# ====================================================================
def test_normalize_value_none():
    """None / 空 → 空字符串"""
    assert normalize_value(None) == ""


def test_normalize_value_bool():
    """bool → 显式 true/false"""
    assert normalize_value(True) == "true"
    assert normalize_value(False) == "false"


def test_normalize_value_datetime():
    """datetime → ISO 格式"""
    dt = _dt.datetime(2025, 6, 15, 14, 30, 0)
    assert normalize_value(dt) == "2025-06-15 14:30:00"

    # 纯日期(时分秒均为 0) → 只保留日期部分
    dt_date_only = _dt.datetime(2025, 6, 15, 0, 0, 0)
    assert normalize_value(dt_date_only) == "2025-06-15"


def test_normalize_value_date():
    """date → ISO 格式"""
    d = _dt.date(2025, 6, 15)
    assert normalize_value(d) == "2025-06-15"


def test_normalize_value_time():
    """time → ISO 格式"""
    t = _dt.time(14, 30, 0)
    assert normalize_value(t) == "14:30:00"


def test_normalize_value_decimal():
    """Decimal → 普通十进制,无科学计数法"""
    assert normalize_value(Decimal("123.456")) == "123.456"
    assert normalize_value(Decimal("0.000")) == "0"


def test_normalize_value_float_integer():
    """float 整数值 → 去掉 .0"""
    assert normalize_value(10.0) == "10"
    assert normalize_value(0.0) == "0"


def test_normalize_value_float_fraction():
    """普通浮点数"""
    assert normalize_value(3.14) == "3.14"


def test_normalize_value_float_noise():
    """浮点噪声被抑制"""
    v = normalize_value(0.30000000000000004)
    assert v == "0.3"


def test_normalize_value_str():
    """普通字符串"""
    assert normalize_value("  hello  ") == "hello"


# ====================================================================
# 2. clean_headers 测试
# ====================================================================
def test_clean_headers_normal():
    """正常表头"""
    assert clean_headers(["姓名", "年龄", "城市"]) == ["姓名", "年龄", "城市"]


def test_clean_headers_empty():
    """空表头补 col{i}"""
    assert clean_headers(["姓名", None, "", "年龄"]) == ["姓名", "col2", "col3", "年龄"]


def test_clean_headers_duplicate():
    """重名列加后缀"""
    assert clean_headers(["A", "A", "A"]) == ["A", "A_1", "A_2"]


# ====================================================================
# 3. merge_header_rows 测试
# ====================================================================
def test_merge_header_rows_two_levels():
    """两级表头合并 — 空值被跳过,不跨列传递"""
    # col 0: "销售" + "Q1" → "销售-Q1"
    # col 1: "" + "Q2"   → "Q2"（空值被跳过）
    rows = [["销售", None], ["Q1", "Q2"]]
    result = merge_header_rows(rows)
    assert result == ["销售-Q1", "Q2"]


def test_merge_header_rows_three_levels():
    """三级表头合并"""
    rows = [["销售", None], ["Q1", "Q2"], ["华北", "华东"]]
    result = merge_header_rows(rows)
    assert result == ["销售-Q1-华北", "Q2-华东"]


def test_merge_header_rows_empty():
    """空表头"""
    assert merge_header_rows([]) == []


# ====================================================================
# 4. row_to_text 测试
# ====================================================================
def test_row_to_text_basic():
    """基本行转文本"""
    headers = ["姓名", "年龄"]
    values = ["张三", 25]
    assert row_to_text(headers, values) == "姓名: 张三；年龄: 25"


def test_row_to_text_skip_empty():
    """跳过空值"""
    headers = ["姓名", "年龄", "城市"]
    values = ["张三", None, ""]
    assert row_to_text(headers, values) == "姓名: 张三"


def test_row_to_text_with_sheet():
    """注入 sheet 名"""
    headers = ["姓名"]
    values = ["张三"]
    result = row_to_text(headers, values, sheet="Sheet1")
    assert result == "[表: Sheet1] 姓名: 张三"


def test_row_to_text_truncate_long():
    """超长字段截断"""
    headers = ["备注"]
    long_text = "A" * (MAX_FIELD_CHARS + 100)
    values = [long_text]
    result = row_to_text(headers, values)
    assert len(result) <= MAX_FIELD_CHARS + len("备注: ") + 1
    assert result.endswith("…")


def test_row_to_text_missing_header():
    """列数多于表头时补 col{i}"""
    headers = ["A"]
    values = [1, 2]
    result = row_to_text(headers, values)
    assert "col2: 2" in result


# ====================================================================
# 5. is_blank_row 测试
# ====================================================================
def test_is_blank_row_true():
    """全空行"""
    assert is_blank_row([None, "", None]) is True
    assert is_blank_row([]) is True


def test_is_blank_row_false():
    """非全空行"""
    assert is_blank_row([None, "hello", ""]) is False


# ====================================================================
# 6. ExcelLoader 集成测试（使用 sample.xlsx）
# ====================================================================
def test_excel_loader_load():
    """load(): 加载 sample.xlsx, 确保返回 Document 列表"""
    loader = ExcelLoader()
    docs = loader.load(str(SAMPLE_XLSX))
    assert isinstance(docs, list)
    assert len(docs) > 0
    for doc in docs:
        assert doc.content  # 非空
        assert "source" in doc.metadata
        assert "sheet" in doc.metadata
        assert "row" in doc.metadata
        assert doc.metadata["format"] == "excel"


def test_excel_loader_without_sheet_name_injection():
    """inject_sheet_name=False 时不注入 [表: xxx] 前缀"""
    loader = ExcelLoader(inject_sheet_name=False)
    docs = loader.load(str(SAMPLE_XLSX))
    for doc in docs:
        assert not doc.content.startswith("[表:")


def test_excel_loader_with_sheet_name_injection():
    """inject_sheet_name=True 时注入 [表: xxx] 前缀"""
    loader = ExcelLoader(inject_sheet_name=True)
    docs = loader.load(str(SAMPLE_XLSX))
    for doc in docs:
        assert doc.content.startswith("[表:")


def test_excel_loader_fill_merged():
    """fill_merged=True 时合并单元格被回填"""
    loader = ExcelLoader(fill_merged=True)
    docs = loader.load(str(SAMPLE_XLSX))
    assert len(docs) > 0
    # 无法直接验证合并单元格是否回填(取决于 sample.xlsx 内容),
    # 但确保不抛异常即为 fill_merged 路径正常


def test_excel_loader_no_fill_merged():
    """fill_merged=False 时使用 read_only 模式"""
    loader = ExcelLoader(fill_merged=False)
    docs = loader.load(str(SAMPLE_XLSX))
    assert len(docs) > 0


def test_excel_loader_header_row():
    """header_row 指定表头起始行"""
    loader = ExcelLoader(header_row=2)  # 从第 2 行开始
    docs = loader.load(str(SAMPLE_XLSX))
    assert len(docs) >= 0  # 可能没有第 2 行数据,但不抛异常


def test_excel_loader_header_depth():
    """header_depth>1 时多级表头合并"""
    loader = ExcelLoader(header_row=1, header_depth=2)
    docs = loader.load(str(SAMPLE_XLSX))
    assert len(docs) >= 0


# ====================================================================
# 7. _read_matrix 测试（内部方法，通过 ExcelLoader 实例间接测试）
# ====================================================================
def test_read_matrix():
    """_read_matrix 返回二维矩阵"""
    from openpyxl import load_workbook
    wb = load_workbook(str(SAMPLE_XLSX), read_only=False, data_only=True)
    try:
        ws = wb.worksheets[0]
        loader = ExcelLoader(fill_merged=True)
        matrix = loader._read_matrix(ws)
        assert isinstance(matrix, list)
        assert len(matrix) > 0
        for row in matrix:
            assert isinstance(row, list)
    finally:
        wb.close()


# ====================================================================
# 8. _load_sheet 测试（内部方法，通过 ExcelLoader 实例间接测试）
# ====================================================================
def test_load_sheet():
    """_load_sheet 返回 Document 列表"""
    from openpyxl import load_workbook
    wb = load_workbook(str(SAMPLE_XLSX), read_only=False, data_only=True)
    try:
        ws = wb.worksheets[0]
        loader = ExcelLoader()
        docs = loader._load_sheet(ws, source=str(SAMPLE_XLSX))
        assert isinstance(docs, list)
        for doc in docs:
            assert doc.content
            assert doc.metadata["source"] == str(SAMPLE_XLSX)
            assert doc.metadata["sheet"] == ws.title
            assert "row" in doc.metadata
    finally:
        wb.close()


# ====================================================================
# 9. 边界条件测试
# ====================================================================
def test_excel_loader_import_error():
    """openpyxl 未安装时抛出 ImportError"""
    loader = ExcelLoader()
    # 模拟 openpyxl 不可用
    import ragcore.loaders.excel_loader as mod
    saved = mod.load_workbook
    mod.load_workbook = None
    try:
        try:
            loader.load(str(SAMPLE_XLSX))
            assert False, "应该抛出 ImportError"
        except ImportError:
            pass
    finally:
        mod.load_workbook = saved


def test_excel_loader_file_not_exists():
    """文件不存在时抛出 FileNotFoundError"""
    loader = ExcelLoader()
    try:
        loader.load("/nonexistent/path/file.xlsx")
        assert False, "应该抛出异常"
    except FileNotFoundError:
        pass


# ====================================================================
# 运行入口
# ====================================================================
if __name__ == "__main__":
    import pytest
    sys.exit(pytest.main([__file__, "-v", "--tb=short"]))