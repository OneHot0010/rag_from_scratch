"""逐个优化点测试 —— Excel 加载器。

每个测试函数独立可运行，打印中间结果，方便 debug。
运行方式: PYTHONPATH=. python3 test/test_excel.py
"""

from pathlib import Path

from ragcore.loaders.excel_loader import ExcelLoader

# ── 测试用文件路径 ──
SAMPLE_XLSX = Path(__file__).resolve().parent.parent / "sample.xlsx"


# ══════════════════════════════════════════════════════════════════════
#  优化点 1: 合并单元格回填 (fill_merged)
# ══════════════════════════════════════════════════════════════════════

def test_merge_cells():
    """测试合并单元格回填。

    核心逻辑(_read_matrix):
      openpyxl 只在合并区域的左上角返回值，其余单元格为 None。
      fill_merged=True 时，把左上角的值回填到整个合并区域。

    对比 fill_merged=True vs False 的矩阵和 Document 输出。
    """
    print("=" * 60)
    print("测试: 合并单元格回填 (fill_merged)")
    print("=" * 60)

    # ── 1. 底层矩阵: fill_merged=True ──
    print("\n[1] fill_merged=True —— 原始矩阵（回填后）:")
    loader_fill = ExcelLoader(header_row=1, fill_merged=True)
    import openpyxl
    wb = openpyxl.load_workbook(SAMPLE_XLSX)
    for ws in wb.worksheets:
        print(f"\n  Sheet: {ws.title}")
        matrix = loader_fill._read_matrix(ws)
        for i, row in enumerate(matrix):
            print(f"    行{i + 1}: {row}")
    wb.close()

    # ── 2. 底层矩阵: fill_merged=False ──
    print("\n[2] fill_merged=False —— 原始矩阵（不回填，合并区域仅左上角有值）:")
    loader_no_fill = ExcelLoader(header_row=1, fill_merged=False)
    wb = openpyxl.load_workbook(SAMPLE_XLSX)
    for ws in wb.worksheets:
        print(f"\n  Sheet: {ws.title}")
        matrix_no_fill = loader_no_fill._read_matrix(ws)
        for i, row in enumerate(matrix_no_fill):
            print(f"    行{i + 1}: {row}")
    wb.close()

    # ── 3. Document 输出: fill_merged=True ──
    print("\n[3] fill_merged=True —— Document 输出:")
    docs = loader_fill.load(str(SAMPLE_XLSX))
    for doc in docs:
        print(f"  row={doc.metadata['row']}, sheet={doc.metadata['sheet']}: {doc.content}")

    # ── 4. Document 输出: fill_merged=False ──
    print("\n[4] fill_merged=False —— Document 输出:")
    docs_no_fill = loader_no_fill.load(str(SAMPLE_XLSX))
    for doc in docs_no_fill:
        print(f"  row={doc.metadata['row']}, sheet={doc.metadata['sheet']}: {doc.content}")

    print("\n" + "=" * 60)
    print("合并单元格测试完成")
    print("=" * 60)


if __name__ == "__main__":
    test_merge_cells()